"""Create a portable Excel workbook with static KPIs and native charts."""
from datetime import date
from decimal import Decimal
from io import BytesIO

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill

from .cleaner import ProcessingResult

NAVY = "183340"
TEAL = "147D73"
CURRENCY = '"$"#,##0.00;[Red]("$"#,##0.00)'


def write_row(ws, values):
    ws.append(list(values))
    for cell in ws[ws.max_row]:
        # Untrusted CSV text must remain literal, even if it starts with '='.
        if isinstance(cell.value, str):
            cell.data_type = "s"
        elif isinstance(cell.value, (date,)):
            cell.number_format = "yyyy-mm-dd"


def header(ws, row: int):
    for cell in ws[row]:
        if cell.value is not None:
            cell.font = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
            cell.fill = PatternFill("solid", fgColor=NAVY)
            cell.alignment = Alignment(vertical="center")
    ws.row_dimensions[row].height = 26


def finish_table(ws):
    header(ws, 1)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    ws.sheet_view.showGridLines = False
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name="Calibri", size=11, color=NAVY)
            cell.alignment = Alignment(vertical="top")
            if cell.row % 2 == 0:
                cell.fill = PatternFill("solid", fgColor="F2F6F7")
    for column in ws.columns:
        width = max(len(str(cell.value or "")) for cell in column) + 3
        ws.column_dimensions[column[0].column_letter].width = min(max(width, 14), 52)


def build_report(result: ProcessingResult) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Cleaned Data"
    write_row(ws, result.cleaned.columns)
    for row in result.cleaned.itertuples(index=False, name=None):
        write_row(ws, row)
    finish_table(ws)
    for row in ws.iter_rows(min_row=2):
        for index in (7, 8):
            row[index].number_format = CURRENCY

    summary = wb.create_sheet("Summary")
    write_row(summary, ["AUTOMATED SALES REPORT", "USD"])
    write_row(summary, ["Self-Initiated Sample Project | Demo data is fictional"])
    write_row(summary, ["Metric", "Value"])
    for label, key in [("Files", "files"), ("Processed Files", "processed_files"),
                       ("Failed Files", "failed_files"), ("Input Rows", "input_rows"),
                       ("Output Rows", "output_rows"), ("Duplicates Removed", "duplicates_removed"),
                       ("Invalid Rows", "invalid_rows"), ("Missing Values", "missing_values"),
                       ("Total Revenue", "total_revenue"), ("Average Order Value", "average_order_value")]:
        write_row(summary, [label, result.metrics[key]])
        if isinstance(result.metrics[key], Decimal):
            summary.cell(summary.max_row, 2).number_format = CURRENCY
    summary.append([])
    write_row(summary, ["Input Rows = Output Rows + Duplicates Removed + Invalid Rows"])
    write_row(summary, ["Input counts exclude files rejected before row processing."])
    write_row(summary, ["Missing Values counts blank canonical cells before filling/deduplication."])
    write_row(summary, ["First valid order ID wins in upload order; review conflicting duplicates."])
    write_row(summary, ["Supplied total is used for revenue; mismatches are flagged for review."])
    write_row(summary, ["Dates: US MM/DD/YYYY. Money: USD only. Each row represents one order."])
    write_row(summary, ["KPIs are computed values; they do not recalculate if you edit this workbook."])
    summary.merge_cells("A2:P2")
    for index in range(15, 22):
        summary.merge_cells(start_row=index, start_column=1, end_row=index, end_column=3)
        summary.cell(index, 1).alignment = Alignment(wrap_text=True, vertical="center")
        summary.cell(index, 1).font = Font(name="Calibri", size=10, color=NAVY)
        summary.row_dimensions[index].height = 30
    for title, group, anchor in [("Revenue by Source", result.by_source, "E3"),
                                  ("Revenue by Category", result.by_category, "E19")]:
        summary.append([])
        write_row(summary, [title, "Orders", "Revenue (USD)"])
        start = summary.max_row
        header(summary, start)
        for row in group.itertuples(index=False, name=None):
            write_row(summary, row)
            summary.cell(summary.max_row, 3).number_format = CURRENCY
        if not group.empty:
            chart = BarChart()
            chart.title = title
            chart.y_axis.title = "Revenue (USD)"
            chart.x_axis.title = "Source" if "Source" in title else "Category"
            chart.add_data(Reference(summary, min_col=3, min_row=start, max_row=summary.max_row), titles_from_data=True)
            chart.set_categories(Reference(summary, min_col=1, min_row=start + 1, max_row=summary.max_row))
            chart.style = 10
            chart.width, chart.height = 21, 8
            chart.legend = None
            summary.add_chart(chart, anchor)
    summary.freeze_panes = "B4"
    summary.sheet_view.showGridLines = False
    for col, width in {"A": 40, "B": 19, "C": 20, "D": 3}.items():
        summary.column_dimensions[col].width = width
    header(summary, 1)
    header(summary, 3)
    summary.row_dimensions[1].height = 32
    summary.sheet_properties.pageSetUpPr.fitToPage = True
    summary.page_setup.orientation = "landscape"
    summary.page_setup.paperSize = summary.PAPERSIZE_A3
    summary.page_setup.fitToWidth = 1
    summary.page_setup.fitToHeight = 1
    summary.print_options.horizontalCentered = True
    summary.print_area = f"A1:P{max(38, summary.max_row)}"

    quality = wb.create_sheet("Data Quality")
    write_row(quality, result.quality.columns)
    for row in result.quality.itertuples(index=False, name=None):
        write_row(quality, [None if value is None or isinstance(value, float) and value != value else value for value in row])
    finish_table(quality)
    for col in ("F", "G", "H"):
        quality.column_dimensions[col].width = 65
    wb.active = 1
    output = BytesIO()
    wb.save(output)
    return output.getvalue()
