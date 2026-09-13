from io import BytesIO
from datetime import datetime
from openpyxl import load_workbook
from src.cleaner import process_files
from src.reporter import build_report


def test_workbook_structure_styles_and_charts(sample_files):
    wb = load_workbook(BytesIO(build_report(process_files(sample_files))))
    assert wb.sheetnames == ["Cleaned Data", "Summary", "Data Quality"]
    clean = wb["Cleaned Data"]
    assert clean.max_row == 16
    assert clean.max_column == 10
    assert clean["A1"].font.bold
    assert clean.freeze_panes == "A2"
    assert clean.auto_filter.ref == "A1:J16"
    assert isinstance(clean["B2"].value, datetime)
    assert clean["B2"].number_format == "yyyy-mm-dd"
    assert clean["I2"].value == 1299
    assert "$" in clean["I2"].number_format
    summary = wb["Summary"]
    assert summary["B12"].value == 4545
    assert summary["B13"].value == 303
    assert len(summary._charts) == 2
    for chart in summary._charts:
        assert chart.series[0].val.numRef.f.startswith("'Summary'!")
    assert wb["Data Quality"].freeze_panes == "A2"


def test_untrusted_formula_text_stays_literal(csv_file):
    r = process_files([csv_file(customer_name='=HYPERLINK("https://example.com")',
                                product="=1+1", customer_email="@SUM(A1)")])
    wb = load_workbook(BytesIO(build_report(r)))
    assert wb["Cleaned Data"]["E2"].value == "=1+1"
    for ws in wb:
        assert all(cell.data_type != "f" for row in ws for cell in row)
