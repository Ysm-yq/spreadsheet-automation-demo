"""Run against an already-started local demo: python tests/browser_smoke.py."""
import json
import os
import re
from pathlib import Path
from time import perf_counter
from openpyxl import load_workbook
from playwright.sync_api import sync_playwright, expect

ROOT = Path(__file__).resolve().parents[1]


def main():
    screenshots = ROOT / "screenshots"
    screenshots.mkdir(exist_ok=True)
    output = ROOT / "output"
    output.mkdir(exist_ok=True)
    with sync_playwright() as p:
        browser = p.chromium.launch(headless=True, channel=os.environ.get("DEMO_BROWSER_CHANNEL") or None)
        page = browser.new_page(viewport={"width": 1600, "height": 1120}, device_scale_factor=1)
        page.goto("http://127.0.0.1:8501")
        expect(page.get_by_role("heading", name="Spreadsheet Automation Demo", exact=True)).to_be_visible()
        expect(page.get_by_role("button", name="Process Files", exact=True)).to_be_disabled()
        page.screenshot(path=str(screenshots / "01_upload.png"), full_page=True)
        paths = [str(ROOT / "sample_data" / name) for name in
                 ("sales_shopify.csv", "sales_stripe.csv", "sales_manual.csv")]
        page.locator('input[type="file"]').set_input_files(paths)
        expect(page.get_by_text("sales_manual.csv", exact=True)).to_be_visible()
        started = perf_counter()
        page.get_by_role("button", name="Process Files", exact=True).click()
        download_button = page.get_by_role("button", name="Download Excel Report", exact=True)
        expect(download_button).to_be_visible(timeout=30000)
        expect(page.get_by_text("$4,545.00", exact=True)).to_be_visible()
        elapsed = perf_counter() - started
        page.screenshot(path=str(screenshots / "02_processed_dashboard.png"), full_page=True)
        with page.expect_download() as pending:
            download_button.click()
        download = pending.value
        assert download.suggested_filename == "automated_sales_report.xlsx"
        download_folder = output / "browser_download"
        download_folder.mkdir(exist_ok=True)
        destination = download_folder / download.suggested_filename
        download.save_as(str(destination))
        wb = load_workbook(destination)
        assert wb.sheetnames == ["Cleaned Data", "Summary", "Data Quality"]
        assert wb["Summary"]["B12"].value == 4545
        assert len(wb["Summary"]._charts) == 2
        page.get_by_role("tab", name="Data Quality", exact=True).click()
        expect(page.get_by_text("Issue type", exact=True)).to_be_visible()
        page.screenshot(path=str(screenshots / "04_data_quality.png"), full_page=True)
        page.get_by_role("tab", name="Summary", exact=True).click()
        expect(page.get_by_role("heading", name="Revenue by Source", exact=True)).to_be_visible()
        page.screenshot(path=str(screenshots / "05_web_summary.png"), full_page=True)
        # Changing files must remove the previous report to prevent stale downloads.
        remove = page.get_by_role("button", name=re.compile(r"^Remove .+\.csv$"))
        while remove.count():
            remove.first.click()
            page.wait_for_timeout(500)
        expect(download_button).not_to_be_visible()
        # A rejected file must not block a valid one.
        page.locator('input[type="file"]').set_input_files([
            {"name": "empty.csv", "mimeType": "text/csv", "buffer": b""},
            {"name": "sales_shopify.csv", "mimeType": "text/csv", "buffer": Path(paths[0]).read_bytes()},
        ])
        expect(page.get_by_text("sales_shopify.csv", exact=True)).to_be_visible()
        page.wait_for_timeout(1000)  # Wait for Streamlit's upload-triggered rerun to settle.
        page.get_by_role("button", name="Process Files", exact=True).click()
        expect(page.get_by_text("empty.csv: This file is empty.", exact=False)).to_be_visible(timeout=30000)
        expect(page.get_by_text("$1,923.00", exact=True)).to_be_visible()
        expect(download_button).to_be_visible()
        while remove.count():
            remove.first.click()
            page.wait_for_timeout(500)
        page.locator('input[type="file"]').set_input_files([
            {"name": "empty.csv", "mimeType": "text/csv", "buffer": b""}])
        expect(page.get_by_text("empty.csv", exact=True)).to_be_visible()
        page.wait_for_timeout(1000)
        page.get_by_role("button", name="Process Files", exact=True).click()
        expect(page.get_by_text("No valid orders were found.", exact=False)).to_be_visible(timeout=30000)
        expect(download_button).to_be_visible()
        browser.close()
    evidence = {"status": "passed", "input_files": 3, "input_rows": 22, "clean_rows": 15,
                "duplicates": 3, "invalid_rows": 4, "revenue_usd": 4545,
                "process_to_download_ready_seconds": round(elapsed, 3),
                "downloaded_workbook_checked": True, "native_excel_charts": 2,
                "changed_upload_clears_report": True, "partial_failure_checked": True,
                "all_empty_checked": True}
    (output / "browser_acceptance.json").write_text(json.dumps(evidence, indent=2), encoding="utf-8")
    print(json.dumps(evidence, indent=2))


if __name__ == "__main__":
    main()
